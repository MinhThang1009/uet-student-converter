from pathlib import Path

from openpyxl import Workbook, load_workbook

from uet_student_converter.converter import parse_page_text, write_excel
from uet_student_converter.models import StudentRecord


def test_parse_page_text_extracts_context_and_records() -> None:
    text = """
    DANH SÁCH SINH VIÊN LỚP QH-2025-I/CQ-X-X1 (K70X-X1)
    Ngành Khoa học dữ liệu
    STT
    MSSV
    Họ và tên
    Ngày sinh
    Giới tính
    1
    25000001
    Nguyễn Sinh Viên
    03/01/2007
    Nam
    2
    25000002
    Trần Thị Mẫu
    10/12/2006
    Nữ
    1
    """

    class_code, major, records = parse_page_text(text)

    assert class_code == "K70X-X1"
    assert major == "Khoa học dữ liệu"
    assert len(records) == 2
    assert records[0].mssv == "25000001"
    assert records[1].password == "10122006"


def test_parse_page_text_uses_previous_context_for_continuation_page() -> None:
    text = """
    STT
    MSSV
    Họ và tên
    Ngày sinh
    Giới tính
    3
    25000003
    Lê Trang Test
    01/02/2007
    Nữ
    2
    """

    class_code, major, records = parse_page_text(
        text,
        current_class="K70X-X1",
        current_major="Khoa học dữ liệu",
    )

    assert class_code == "K70X-X1"
    assert major == "Khoa học dữ liệu"
    assert records[0].source_stt == 3


def test_write_excel_preserves_leading_zero_password(tmp_path: Path) -> None:
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "output.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "K70"
    for column, header in enumerate(
        ("STT", "Họ và tên", "Giới tính", "Ngành", "Tài khoản", "Mật khẩu", "Tình trạng"),
        start=2,
    ):
        worksheet.cell(row=2, column=column, value=header)
    workbook.save(template_path)

    records = [
        StudentRecord(
            source_stt=1,
            mssv="25000001",
            name="Nguyễn Sinh Viên",
            date_of_birth="03/01/2007",
            gender="Nam",
            major="Khoa học dữ liệu",
            class_code="K70X-X1",
        )
    ]

    assert write_excel(records, template_path, output_path) == 1

    result = load_workbook(output_path, data_only=False)
    worksheet = result["K70"]
    assert worksheet["B3"].value == 1
    assert worksheet["F3"].value == "25000001@vnu.edu.vn"
    assert worksheet["G3"].value == "03012007"
    assert worksheet["G3"].number_format == "@"
    assert worksheet["F3"].hyperlink.target == "mailto:25000001@vnu.edu.vn"
