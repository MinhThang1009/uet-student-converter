"""Extract students from PDF and write them using the UET Excel template."""

from __future__ import annotations

import re
from collections.abc import Sequence
from pathlib import Path

import pymupdf
from openpyxl import load_workbook

from .models import StudentRecord

TITLE_PATTERN = re.compile(r"DANH SÁCH SINH VIÊN LỚP .*?\((K70[^)]+)\)")
MAJOR_PATTERN = re.compile(r"^\s*Ngành\s+(.+?)\s*$", re.MULTILINE)
STT_PATTERN = re.compile(r"^\d+$")
MSSV_PATTERN = re.compile(r"^\d{8}$")
DATE_PATTERN = re.compile(r"^\d{2}/\d{2}/\d{4}$")
GENDERS = {"Nam", "Nữ"}

HEADERS = (
    "STT",
    "Họ và tên",
    "Giới tính",
    "Ngành",
    "Tài khoản",
    "Mật khẩu",
    "Tình trạng",
)


def _page_text(page: pymupdf.Page) -> str:
    """Return text from one PDF page in reading order."""

    return "\n".join(
        block[4] for block in page.get_text("blocks", sort=True) if len(block) > 4 and block[4]
    )


def parse_page_text(
    text: str,
    *,
    current_class: str | None = None,
    current_major: str | None = None,
) -> tuple[str | None, str | None, list[StudentRecord]]:
    """Parse one page and return the updated context and extracted records."""

    title_match = TITLE_PATTERN.search(text)
    if title_match:
        current_class = title_match.group(1)

    major_match = MAJOR_PATTERN.search(text)
    if major_match:
        current_major = major_match.group(1).strip()

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    records: list[StudentRecord] = []
    for index in range(len(lines) - 4):
        stt, mssv, name, date_of_birth, gender = lines[index : index + 5]
        if not STT_PATTERN.fullmatch(stt):
            continue
        if not MSSV_PATTERN.fullmatch(mssv):
            continue
        if not DATE_PATTERN.fullmatch(date_of_birth) or gender not in GENDERS:
            continue
        if current_class is None or current_major is None:
            raise ValueError("Could not determine the class or major before a student record")
        records.append(
            StudentRecord(
                source_stt=int(stt),
                mssv=mssv,
                name=name,
                date_of_birth=date_of_birth,
                gender=gender,
                major=current_major,
                class_code=current_class,
            )
        )
    return current_class, current_major, records


def extract_students(pdf_path: Path) -> list[StudentRecord]:
    """Read all students from the PDF in source order."""

    if not pdf_path.is_file():
        raise FileNotFoundError(f"Source PDF was not found: {pdf_path}")

    records: list[StudentRecord] = []
    current_class: str | None = None
    current_major: str | None = None
    with pymupdf.open(pdf_path) as document:
        for page in document:
            current_class, current_major, page_records = parse_page_text(
                _page_text(page),
                current_class=current_class,
                current_major=current_major,
            )
            records.extend(page_records)

    if not records:
        raise ValueError(f"No student records could be extracted from: {pdf_path}")
    _validate_records(records)
    return records


def _validate_records(records: Sequence[StudentRecord]) -> None:
    """Validate the data conditions required before writing Excel."""

    mssv_values = [record.mssv for record in records]
    if len(mssv_values) != len(set(mssv_values)):
        raise ValueError("The PDF contains duplicate MSSV values")
    if any(len(record.password) != 8 for record in records):
        raise ValueError("A date-of-birth password is not eight characters long")


def write_excel(
    records: Sequence[StudentRecord],
    template_path: Path,
    output_path: Path,
) -> int:
    """Write records to an Excel workbook using the template structure."""

    if not records:
        raise ValueError("There are no records to write to Excel")
    if not template_path.is_file():
        raise FileNotFoundError(f"Excel template was not found: {template_path}")

    _validate_records(records)
    workbook = load_workbook(template_path)
    if "K70" not in workbook.sheetnames:
        raise ValueError("The Excel template does not contain a K70 sheet")
    worksheet = workbook["K70"]

    if worksheet.max_row >= 3:
        worksheet.delete_rows(3, worksheet.max_row - 2)
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.hyperlink is not None:
                cell.hyperlink = None

    for column, header in enumerate(HEADERS, start=2):
        worksheet.cell(row=2, column=column, value=header)

    for output_stt, record in enumerate(records, start=1):
        row = output_stt + 2
        values = (
            output_stt,
            record.name,
            record.gender,
            record.major,
            record.account,
            record.password,
            "Hoạt động",
        )
        for column, value in enumerate(values, start=2):
            cell = worksheet.cell(row=row, column=column, value=value)
            if column == 6:
                cell.hyperlink = f"mailto:{record.account}"
                cell.style = "Hyperlink"
            elif column == 7:
                cell.number_format = "@"

    worksheet.auto_filter.ref = f"B2:H{len(records) + 2}"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return len(records)


def convert_pdf_to_excel(
    pdf_path: Path,
    template_path: Path,
    output_path: Path,
) -> int:
    """Convert a PDF to Excel and return the number of written records."""

    records = extract_students(pdf_path)
    return write_excel(records, template_path, output_path)
